"""Replay Scout - enter Showdown username(s), get a scouting workbook.
Run:  pip install -r requirements.txt && python app.py   ->  http://localhost:5000
"""
import io, itertools, json, os, re, threading, time, uuid
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor

import requests
from flask import Flask, jsonify, render_template_string, request, send_file
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = "https://replay.pokemonshowdown.com"
HDR = {"User-Agent": "replay-scout/1.0 (personal scouting tool)"}
HERE = os.path.dirname(os.path.abspath(__file__))
CACHE, OUT = os.path.join(HERE, "cache"), os.path.join(HERE, "output")
os.makedirs(CACHE, exist_ok=True); os.makedirs(OUT, exist_ok=True)
norm = lambda s: re.sub(r"[^a-z0-9]", "", (s or "").lower())

# ---------------------------------------------------------------- fetching
def search(name, fmt, max_pages=40):
    rows = []
    for page in range(1, max_pages + 1):
        p = {"user": name, "page": page}
        if fmt: p["format"] = fmt
        r = requests.get(f"{BASE}/search.json", params=p, headers=HDR, timeout=20)
        r.raise_for_status()
        batch = r.json()
        rows += batch
        if len(batch) < 50: break
        time.sleep(0.3)
    return rows

def get_replay(rid):
    f = os.path.join(CACHE, rid + ".json")
    if os.path.exists(f):
        with open(f) as fh: return json.load(fh)
    r = requests.get(f"{BASE}/{rid}.json", headers=HDR, timeout=20)
    r.raise_for_status()
    d = r.json()
    with open(f, "w") as fh: json.dump(d, fh)
    time.sleep(0.1)
    return d

# ----------------------------------------------------------------- sprites
SPR = os.path.join(CACHE, "sprites"); os.makedirs(SPR, exist_ok=True)
PS = "https://play.pokemonshowdown.com/sprites"
KEEP_HYPHEN = {"porygon-z": "porygonz", "ho-oh": "hooh", "jangmo-o": "jangmoo", "hakamo-o": "hakamoo", "kommo-o": "kommoo"}

def sprite_ids(species):
    low = species.lower()
    for k, v in KEEP_HYPHEN.items():
        if low.startswith(k): low = v + low[len(k):]; break
    base, _, forme = low.partition("-")
    a = lambda t: re.sub(r"[^a-z0-9]", "", t)
    ids = [a(base) + ("-" + a(forme) if forme else "")]
    if forme: ids.append(a(base))
    return ids

def fetch_sprite(species):
    """Download a small PNG icon for a species (cached). Returns file path or None."""
    out = os.path.join(SPR, re.sub(r"[^a-z0-9-]", "", species.lower()) + ".png")
    if os.path.exists(out): return out
    for sid in sprite_ids(species):
        for folder in ("xy", "gen5", "dex"):
            try:
                r = requests.get(f"{PS}/{folder}/{sid}.png", headers=HDR, timeout=15)
                if r.status_code == 200 and r.content[:4] == b"\x89PNG":
                    im = PILImage.open(io.BytesIO(r.content)).convert("RGBA")
                    im.thumbnail((48, 48)); im.save(out)
                    return out
            except Exception:
                pass
    return None

# ----------------------------------------------------------------- parsing
def parse(log):
    teams = {"p1": [], "p2": []}
    nick = {"p1": {}, "p2": {}}
    out = {"names": {}, "turns": 0, "winner": None, "lead": {}}

    def who(tok):
        m = re.match(r"(p[12])[a-d]?: (.*)", tok or "")
        return (m.group(1), m.group(2)) if m else (None, None)

    def mon(tok):
        s, n = who(tok)
        return nick[s].get(n) if s else None

    def find_or_add(side, species):
        for m in teams[side]:
            if m["species"] == species and not m["seen"]: return m
        for m in teams[side]:
            if m["species"].endswith("-*") and species.startswith(m["species"][:-1]) and not m["seen"]:
                m["species"] = species; return m
        for m in teams[side]:
            if m["species"] == species: return m
        m = {"species": species, "ability": None, "item": None, "moves": [], "seen": False}
        teams[side].append(m)
        return m

    for line in log.split("\n"):
        if not line.startswith("|"): continue
        p = line.split("|")
        k = p[1] if len(p) > 1 else ""
        if k == "player" and len(p) > 3: out["names"][p[2]] = p[3]
        elif k == "poke":
            teams[p[2]].append({"species": p[3].split(",")[0], "ability": None, "item": None, "moves": [], "seen": False})
        elif k in ("switch", "drag"):
            s, n = who(p[2])
            m = find_or_add(s, p[3].split(",")[0]); m["seen"] = True
            nick[s][n] = m
            out["lead"].setdefault(s, m)
        elif k == "detailschange":
            m = mon(p[2])
            if m: m["species"] = p[3].split(",")[0]
        elif k == "-mega":
            m = mon(p[2])
            if m and len(p) > 4 and p[4]: m["item"] = p[4]
        elif k == "move":
            m = mon(p[2]); tags = "|".join(p[5:])
            if m and p[3] not in ("Struggle", "Recharge") and ("[from]" not in tags or "lockedmove" in tags):
                if p[3] not in m["moves"]: m["moves"].append(p[3])
        elif k == "turn": out["turns"] = int(p[2])
        elif k == "win": out["winner"] = p[2]
        elif k.startswith("-"):
            tags = p[4:]
            frm = next((t for t in tags if t.startswith("[from]")), "")
            of = next((t[5:].strip() for t in tags if t.startswith("[of]")), None)
            target = mon(of) if of else mon(p[2] if len(p) > 2 else "")
            mt = re.match(r"\[from\] (ability|item): (.*)", frm)
            if k == "-ability" and len(p) > 3:
                m = mon(p[2])
                if m:
                    m["ability"] = "Trace" if "Trace" in frm else (p[3] if not mt else m["ability"])
            elif k in ("-item", "-enditem") and len(p) > 3:
                m = mon(p[2])
                if m and not re.search(r"Trick|Switcheroo|Bestow|Thief|Covet", frm): m["item"] = p[3]
            elif k == "-activate" and len(p) > 3 and p[3].startswith("ability: "):
                m = mon(p[2])
                if m: m["ability"] = p[3][9:]
            if mt and target and not (k == "-ability" and "Trace" in frm):
                if mt.group(1) == "ability": target["ability"] = mt.group(2)
                else: target["item"] = mt.group(2)
    out["teams"] = teams
    return out

def to_game(rid, d, norms, found_as):
    p = parse(d.get("log", ""))
    side = next((s for s in ("p1", "p2") if norm(p["names"].get(s)) in norms), None)
    if not side: return None
    opp = "p2" if side == "p1" else "p1"
    me, op = p["names"][side], p["names"].get(opp, "")
    return dict(id=rid, ts=d.get("uploadtime", 0), date=time.strftime("%Y-%m-%d", time.gmtime(d.get("uploadtime", 0))),
                format=d.get("format", ""), turns=p["turns"], me=me, opp=op, found_as=found_as,
                team=p["teams"][side], opp_team=p["teams"][opp],
                lead=p["lead"][side]["species"] if side in p["lead"] else "",
                win=norm(p["winner"]) == norm(me), winner=p["winner"] or "")

# --------------------------------------------------------------- workbook
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)

def put(ws, r, c, v, bold=False, bg=None, fmt=None):
    cell = ws.cell(r, c, v); cell.font = BOLD if bold else FONT
    if bg: cell.fill = fill(bg)
    if fmt: cell.number_format = fmt
    return cell

def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x

def build(games, names, path, sprites=None):
    sprites = sprites or {}
    wb = Workbook()
    def icon(ws, species, r, c):
        p = sprites.get(species)
        if p: ws.add_image(XLImage(p), f"{get_column_letter(c)}{r}")
    # Replay Summary  (icon column + name column for every Pokemon)
    ws = wb.active; ws.title = "Replay Summary"
    heads = ["Replay", "Date", "Format", "Turns", "Scouted Player"]
    for i in range(1, 7): heads += [f"P{i}", f"P{i} Name"]
    heads.append("Opponent")
    for i in range(1, 7): heads += [f"O{i}", f"O{i} Name"]
    heads += ["Result", "Winner"]
    put(ws, 1, 1, f"Replay Summary - {', '.join(names)}", True)
    for i, h in enumerate(heads, 1): put(ws, 2, i, h, True, "DCEBFF")
    for r, g in enumerate(games, 3):
        ws.row_dimensions[r].height = 38
        vals = {1: g["id"], 2: g["date"], 3: g["format"], 4: g["turns"], 5: g["me"], 18: g["opp"],
                31: "W" if g["win"] else "L", 32: g["winner"]}
        for start, team in ((6, g["team"]), (19, g["opp_team"])):
            for i in range(6):
                if i < len(team):
                    vals[start + 2 * i + 1] = team[i]["species"]
                    icon(ws, team[i]["species"], r, start + 2 * i)
        for c in range(1, 33):
            cell = put(ws, r, c, vals.get(c), bg=("E2F0D9" if g["win"] else "FBE5D6") if c == 31 else None)
            cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A3"
    widths(ws, [36, 11, 12, 7, 16] + [7, 15] * 6 + [16] + [7, 15] * 6 + [7, 16])
    # Replay Links
    ws = wb.create_sheet("Replay Links")
    for i, h in enumerate(["Replay URL", "Date", "Format", "Scouted Player", "Found As", "Opponent", "Result", "Winner"], 1):
        put(ws, 1, i, h, True, "DCEBFF")
    for r, g in enumerate(games, 2):
        for c, v in enumerate([f"{BASE}/{g['id']}", g["date"], g["format"], g["me"], g["found_as"], g["opp"],
                               "W" if g["win"] else "L", g["winner"]], 1): put(ws, r, c, v)
    ws.freeze_panes = "A2"; widths(ws, [60, 11, 12, 16, 16, 18, 7, 16])
    # Teams Details
    ws = wb.create_sheet("Teams Details"); r = 1
    for g in games:
        put(ws, r, 1, f"{g['id']}  |  {g['date']}  |  vs {g['opp']}  |  {'W' if g['win'] else 'L'}", True, "92CDDC")
        for c in range(2, 9): put(ws, r, c, None, bg="92CDDC")
        labels = ["Icon", "Pokemon", "Ability", "Item", "Tera", "Move 1", "Move 2", "Move 3", "Move 4"]
        for i, lab in enumerate(labels): put(ws, r + 1 + i, 2, lab, True, "B7DEE8")
        ws.row_dimensions[r + 1].height = 38
        for j in range(6):
            m = g["team"][j] if j < len(g["team"]) else None
            vals = [None, m["species"], m["ability"], m["item"], None] + [(m["moves"][k] if k < len(m["moves"]) else None) for k in range(4)] if m else [None] * 9
            for i, v in enumerate(vals): put(ws, r + 1 + i, 3 + j, v, bold=(i == 1), bg="B7DEE8" if i < 2 else "DDEFF3")
            if m: icon(ws, m["species"], r + 1, 3 + j)
        r += 11
    widths(ws, [4, 10] + [17] * 6)
    # Player Usage Stats
    ws = wb.create_sheet("Player Usage Stats"); N = len(games)
    put(ws, 1, 1, f"{', '.join(names)} Usage Stats  ({N} games)", True)
    def table(col0, title_cols, rows):
        for i, h in enumerate(title_cols + ["Use", "Usage %", "Win %"]): put(ws, 2, col0 + i, h, True, "DCEBFF")
        for r, (key, use, wins) in enumerate(rows, 3):
            for i, v in enumerate(key): put(ws, r, col0 + i, v)
            n = len(key)
            put(ws, r, col0 + n, use); put(ws, r, col0 + n + 1, use / N, fmt="0%"); put(ws, r, col0 + n + 2, wins / use, fmt="0%")
    def tally(keyfn):
        use, wins = Counter(), Counter()
        for g in games:
            for k in keyfn(g): use[k] += 1; wins[k] += g["win"]
        return sorted(((k, use[k], wins[k]) for k in use), key=lambda x: (-x[1], x[0]))
    mons = tally(lambda g: {(m["species"],) for m in g["team"]})
    table(1, ["Pokemon"], mons)
    table(6, ["Lead"], tally(lambda g: [(g["lead"],)] if g["lead"] else [])[:30])
    table(11, ["Core P1", "P2"], [x for x in tally(lambda g: itertools.combinations(sorted({m["species"] for m in g["team"]}), 2)) if x[1] > 1][:40])
    table(17, ["Core P1", "P2", "P3"], [x for x in tally(lambda g: itertools.combinations(sorted({m["species"] for m in g["team"]}), 3)) if x[1] > 1][:40])
    ws.freeze_panes = "A3"; widths(ws, [20, 7, 9, 7, 3, 20, 7, 9, 7, 3, 18, 18, 7, 9, 7, 3, 16, 16, 16, 7, 9, 7])
    wb.save(path)

# ------------------------------------------------------------------- jobs
JOBS = {}

def worker(jid, names, fmt, opps, pasted):
    j = JOBS[jid]
    try:
        seen, per = {}, Counter()
        for n in names:
            j["msg"] = f"Searching Showdown for '{n}'..."
            for row in search(n, fmt):
                if row.get("private"): continue
                if row["id"] not in seen: seen[row["id"]] = n; per[n] += 1
            j["per_name"] = dict(per)
        for rid in pasted:
            if rid not in seen: seen[rid] = "pasted"; per["pasted links"] += 1
        j["per_name"] = dict(per)
        ids = list(seen); j["total"] = len(ids)
        if not ids: raise RuntimeError("No public replays found for: " + ", ".join(names or ["(none)"]) + ". Showdown only knows Showdown usernames - use Saved aliases to link your forum name to your Showdown names, and check the tier.")
        j["msg"] = "Downloading and parsing replays..."
        def task(rid):
            try: return rid, get_replay(rid)
            except Exception: return rid, None
        data, failed = {}, 0
        with ThreadPoolExecutor(4) as ex:
            for rid, d in ex.map(task, ids):
                j["done"] += 1
                if d: data[rid] = d
                else: failed += 1
        if not names:  # only replay links given: scout the name that appears most often
            c = Counter(n for d in data.values() for n in re.findall(r"\|player\|p[12]\|([^|]+)\|", d.get("log", "")))
            if not c: raise RuntimeError("Could not read any of those replays.")
            names = [c.most_common(1)[0][0]]
        norms, onorms = {norm(n) for n in names}, {norm(o) for o in opps}
        games = []
        for rid, d in data.items():
            g = to_game(rid, d, norms, seen[rid])
            if g and (not onorms or norm(g["opp"]) in onorms): games.append(g)
            else: failed += 1
        if not games: raise RuntimeError("Replays were found, but none matched the name/opponent filter.")
        games.sort(key=lambda g: g["ts"], reverse=True)
        j["msg"] = "Fetching sprites..."
        species = {m["species"] for g in games for m in g["team"] + g["opp_team"]}
        with ThreadPoolExecutor(6) as ex:
            sprites = dict(zip(species, ex.map(fetch_sprite, species)))
        build(games, names, os.path.join(OUT, jid + ".xlsx"), sprites)
        j.update(state="done", msg=f"Done - {len(games)} games" + (f" ({failed} skipped)" if failed else ""), games=len(games))
    except Exception as e:
        j.update(state="error", msg=f"{type(e).__name__}: {e}")

app = Flask(__name__)

@app.get("/")
def index(): return render_template_string(PAGE)

@app.post("/start")
def start():
    body = request.get_json(force=True)
    split = lambda t: [n.strip() for n in re.split(r"[,\n]", t or "") if n.strip()]
    names, opps = split(body.get("usernames")), split(body.get("opponents"))
    pasted = list(dict.fromkeys(re.findall(r"(?:replay\.pokemonshowdown\.com/)?((?:[a-z0-9]+-)?gen\d\w*-\d+(?:-[a-z0-9]+pw)?)", (body.get("replays") or "").lower())))
    if not names and not pasted: return jsonify(error="Enter a username or paste replay links"), 400
    jid = uuid.uuid4().hex[:10]
    JOBS[jid] = dict(state="running", msg="Starting...", done=0, total=0, per_name={})
    threading.Thread(target=worker, args=(jid, names, norm(body.get("format", "gen6ou")), opps, pasted), daemon=True).start()
    return jsonify(id=jid)

@app.get("/status/<jid>")
def status(jid): return jsonify(JOBS.get(jid, {"state": "error", "msg": "Unknown job"}))

@app.get("/download/<jid>")
def download(jid):
    return send_file(os.path.join(OUT, jid + ".xlsx"), as_attachment=True, download_name="scouting.xlsx")

PAGE = """<!doctype html><meta name=viewport content="width=device-width,initial-scale=1"><title>Replay Scout</title>
<style>body{font:15px system-ui;max-width:560px;margin:40px auto;padding:0 16px;color:#1f2937}
h1{font-size:22px}textarea,input{width:100%;box-sizing:border-box;padding:10px;border:1px solid #cbd5e1;border-radius:8px;font:inherit;margin:4px 0 14px}
button{background:#2563eb;color:#fff;border:0;border-radius:8px;padding:10px 18px;font:inherit;cursor:pointer}
button:disabled{opacity:.5}.bar{height:8px;background:#e2e8f0;border-radius:4px;overflow:hidden;margin:14px 0}
.bar i{display:block;height:100%;width:0;background:#2563eb;transition:width .3s}small{color:#64748b}a.dl{display:inline-block;margin-top:10px;font-weight:600}</style>
<h1>Replay Scout</h1>
<label>Showdown username(s) <small>- one per line or comma separated. Add alts here too. Leave empty if you paste links.</small></label>
<textarea id=u rows=3 placeholder="White Atoq"></textarea>
<label>Format ID <small>- e.g. gen6ou (blank = all formats)</small></label>
<input id=f value=gen6ou>
<label>Opponents <small>- optional, only games vs these names</small></label>
<input id=o>
<label>Replay links <small>- optional, paste links (one per line)</small></label>
<textarea id=r rows=3></textarea>
<label>Saved aliases <small>- optional. One per line: <b>Sna: Snaurus, altname2, altname3</b>. Then typing "Sna" above searches all of them. Saved on this phone.</small></label>
<textarea id=a rows=3></textarea>
<button id=go>Build scouting sheet</button>
<div class=bar><i id=b></i></div><div id=m></div><div id=p></div>
<script>
const $=i=>document.getElementById(i);
$('a').value=localStorage.getItem('aliases')||'';
const key=t=>t.toLowerCase().replace(/[^a-z0-9]/g,'');
function expand(t){const map={};$('a').value.split('\\n').forEach(l=>{const i=l.indexOf(':');if(i>0)map[key(l.slice(0,i))]=l.slice(i+1)});
  return t.split(/[,\\n]/).map(x=>x.trim()).filter(Boolean).map(n=>map[key(n)]||n).join(',')}
$('go').onclick=async()=>{localStorage.setItem('aliases',$('a').value);
  $('go').disabled=true;$('m').textContent='Starting...';$('p').innerHTML='';$('b').style.width='0';
  const r=await fetch('/start',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({usernames:expand($('u').value),format:$('f').value,opponents:$('o').value,replays:$('r').value})});
  const j=await r.json(); if(j.error){$('m').textContent=j.error;$('go').disabled=false;return}
  const t=setInterval(async()=>{
    const s=await (await fetch('/status/'+j.id)).json();
    $('m').textContent=s.msg+(s.total?` (${s.done}/${s.total})`:'');
    $('b').style.width=(s.total?100*s.done/s.total:5)+'%';
    $('p').innerHTML=Object.entries(s.per_name||{}).map(([k,v])=>`<div><b>${k}</b>: ${v} replays</div>`).join('');
    if(s.state!='running'){clearInterval(t);$('go').disabled=false;
      if(s.state=='done')$('p').innerHTML+=`<a class=dl href="/download/${j.id}">Download scouting.xlsx</a>`}
  },800)};
</script>"""

if __name__ == "__main__":
    app.run(debug=False, host=os.environ.get("HOST", "127.0.0.1"), port=int(os.environ.get("PORT", 5000)))
